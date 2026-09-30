#!/usr/bin/env python3
"""Optional independent reader check using an already installed openpyxl.

Development verification only: this module is not an application, native-build
or required test dependency. It never installs a package or accesses a network.
"""
import json
from pathlib import Path
import sys
import warnings
import xml.etree.ElementTree as ET

import openpyxl

warnings.simplefilter("error", UserWarning)
checks = 0


def check(value, detail):
    global checks
    checks += 1
    if not value:
        raise AssertionError(detail)


directory = Path(sys.argv[1])
themes = ["Classic Light", "Monochromatic Dark", "Corporate Blue", "High Contrast"]
for theme in themes:
    for prefix in ["", "again-"]:
        path = directory / (prefix + theme.replace(" ", "_") + ".xlsx")
        book = openpyxl.load_workbook(path)
        try:
            check(len(book.worksheets) == 1, "one worksheet")
            sheet = book.active
            check(sheet["A2"].value == "مرحبا", "Unicode shared strings")
            check(sheet["B4"].value == "=3/4", "formula source")
            check(sheet["C7"].value == "=SUM(B2:B4)", "aggregate source")
            check(sheet.column_dimensions["A"].width == 18.5, "column width")
            check(sheet.row_dimensions[2].height == 24, "row height")
            check(sheet.freeze_panes == "B1", "frozen column")
            cell = sheet["B2"]
            check(cell.font.name == "Consolas" and cell.font.sz == 14 and cell.font.b and cell.font.i, "font")
            check(cell.font.color.rgb == "FF123456", "font RGB")
            check(cell.fill.fgColor.rgb == "FFABCDEF", "fill RGB")
            check(cell.border.left.style == "double" and cell.border.bottom.style == "medium", "independent borders")
            check(cell.alignment.horizontal == "right" and cell.alignment.vertical == "center" and cell.alignment.wrapText, "alignment/wrap")
            check(cell.number_format == '"USD "#,##0.00', "number format")
            check(ET.fromstring(book.loaded_theme).get("name") == theme, "theme XML")
            check(sheet.tables["Sales"].ref == "A1:B4", "table range")
            check([c.name for c in sheet.tables["Sales"].tableColumns] == ["Label", "Value"], "table headers")
            check(len(sheet._charts) == 3, "chart and drawing relationships")
            for chart, kind in zip(sheet._charts, ["bar", "col", "line"]):
                check(len(chart.series) == 1, "single series")
                check(chart.series[0].cat.strRef.f.endswith("!$A$2:$A$4"), "category range")
                check(chart.series[0].val.numRef.f.endswith("!$B$2:$B$4"), "value range")
                check(chart.series[0].val.numRef.numCache.pt[-1].v == .75, "native value cache")
                check(chart.type == kind if kind != "line" else type(chart).__name__ == "LineChart", "chart kind")
        finally:
            book.close()

book = openpyxl.load_workbook(directory / "totals.xlsx")
try:
    sheet = book.active
    check(sheet["A4"].value == "Total" and sheet["B4"].value == "=SUM(B2:B3)", "native totals formula")
    table = sheet.tables["TotalsTable"]
    check(table.totalsRowCount == 1 and table.tableColumns[1].totalsRowFunction == "sum", "totals metadata")
    check(table.tableColumns[0].totalsRowLabel == "Total", "totals label")
finally:
    book.close()
print(json.dumps({"reader": "openpyxl", "version": openpyxl.__version__, "workbooks": 9, "checks": checks, "failures": 0}))
